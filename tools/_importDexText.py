from openpyxl import load_workbook
from openpyxl.styles import Color, PatternFill, Font, Border

import sys
import openpyxl
import shutil
import os

class bcolors:
    HEADER = '\033[95m'
    OKBLUE = '\033[94m'
    OKCYAN = '\033[96m'
    OKGREEN = '\033[92m'
    WARNING = '\033[93m'
    FAIL = '\033[91m'
    ENDC = '\033[0m'
    BOLD = '\033[1m'
    UNDERLINE = '\033[4m'

def RN(text):
    if text == None:
        return ''
    return text

def FormatDexText(sheet,row,col):
    output = ""
    for i in range(4):
        output += sheet.cell(row=row + i, column = col).value
        output = output.replace('\n','').replace('\r','')
        if i == 1:
            output += '/'
        elif i < 3:
            output += ';'
        else:
            output += '@'
    return output

# Program Start
xlsxInputListPath = sys.argv[1]
xlsxOutputListPath = sys.argv[2]

# Load Workbook
wb = load_workbook(filename = xlsxInputListPath)
wb2 = load_workbook(filename = xlsxOutputListPath)
wb2Sheet = wb2._sheets[0]
filePaths = []

print(bcolors.OKGREEN)
print('正在读取图鉴文本 from: ' + xlsxInputListPath)


sheet = wb._sheets[0]
for dexID in range(0,251):
    row = 3 + 5 * dexID
    goldCol = 3
    silverCol = 13
    goldText = FormatDexText(sheet,row,goldCol)
    silverText = FormatDexText(sheet,row,silverCol)

    wb2Sheet.cell(row=dexID + 1, column = 6).value = goldText
    wb2Sheet.cell(row=dexID + 1, column = 8).value = silverText
    # print(goldText)
    # print(silverText)
    # print()




print('正在将图鉴文本导出到: ' + xlsxOutputListPath)
wb2.save(xlsxOutputListPath)
print(bcolors.OKGREEN)
print('宝可梦图鉴数据已导出到: ' + xlsxOutputListPath)
print()


