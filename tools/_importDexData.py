from openpyxl import load_workbook
from openpyxl.styles import Color, PatternFill, Font, Border

import sys
import openpyxl
import shutil
import os

class bcolors:
    HEADER = '\033[95m'
    OKBLUE = '\033[94m'
    OKCYAN = '\033[96m'
    OKGREEN = '\033[92m'
    WARNING = '\033[93m'
    FAIL = '\033[91m'
    ENDC = '\033[0m'
    BOLD = '\033[1m'
    UNDERLINE = '\033[4m'

def RN(text):
    if text == None:
        return ''
    return text


# Program Start
xlsxListPath = sys.argv[1]


# Load Workbook
wb = load_workbook(filename = xlsxListPath)

filePaths = []

print(bcolors.OKGREEN)
print('正在导入图鉴数据到工程 from: ' + xlsxListPath)




sheet = wb._sheets[0]
for row in range(1,252):
    fileName = sheet.cell(row=row, column = 3).value + ".asm"
    pokemonType = sheet.cell(row=row, column = 4).value
    pokemonHeightWeight = sheet.cell(row=row, column = 11).value
    for ver in range(2):
        outputTextLines = []
        desc = sheet.cell(row=row, column = 6 + 2 * ver).value.replace(';','<NEXT>').replace('/','@')
        filePath = sheet.cell(row=1, column = 1).value + '/' + sheet.cell(row=row, column = 12 + ver).value + '/' + fileName
        outputTextLines.append(f'\tdb_w \"{pokemonType}@\"\n')
        outputTextLines.append(f'\t{pokemonHeightWeight}\n')
        outputTextLines.append(f'\tdb_w \"{desc}\"\n')
        # print(filePath)
        # print(''.join(outputTextLines))
        # print()
        with open(filePath, 'w',encoding='utf-8') as f:
            f.write(''.join(outputTextLines))






print(bcolors.OKGREEN)
print('宝可梦图鉴数据已导入到工程。')
print()


