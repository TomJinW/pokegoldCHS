from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl import load_workbook
import openpyxl
import os
from bs4 import BeautifulSoup
part = '｜'

class bcolors:
    HEADER = '\033[95m'
    OKBLUE = '\033[94m'
    OKCYAN = '\033[96m'
    OKGREEN = '\033[92m'
    WARNING = '\033[93m'
    FAIL = '\033[91m'
    ENDC = '\033[0m'
    BOLD = '\033[1m'
    UNDERLINE = '\033[4m'


def RN(text):
    if text == None:
        return ''
    else:
        return text.strip()
    
def emptyStr(label):
    return RN(label) == ''

def getCrystalDict():
    wb = load_workbook('xlsx/text.xlsx', data_only=True)

    label2EngTextDict = dict()
    for i in range(10):
        ws = wb['文' + str(i + 1)]
        wbRow = 1

        tmpText = []
        replacing = '【NUM】'
        replacements = dict()
        label = ''
        for wbRow in range(1, ws.max_row + 1):
            tmpLabel = ws.cell(row = wbRow, column = 10).value
            tmpcomment = ws.cell(row = wbRow, column = 10).comment
            if not emptyStr(tmpLabel) and tmpcomment is not None:
                if label != '':
                    label2EngTextDict[label] = tmpText
                    # if label == 'torikae_msg_000_D_action':
                    #     print(tmpText)
                label = tmpLabel
                tmpText = []
                replacements = dict()
                row2 = wbRow + 1
                while not emptyStr(ws.cell(row = row2, column = 10).value):
                    tmpRelacing = replacing.replace('NUM',str(row2 - wbRow - 1))
                    replacements[tmpRelacing] = '+' + ws.cell(row = row2, column = 10).value + '+'
                    # print(replacements)
                    row2 += 1
            else:
                tmpLine = RN(ws.cell(row = wbRow, column = 1).value)
                # if label == 'torikae_msg_000_D_action':
                #     print(tmpLine)
                if not emptyStr(tmpLine):
                    if not '结束' in tmpLine:
                        for key in replacements:
                            tmpLine = tmpLine.replace(key,replacements[key])
                            
                        # print(tmpLine)
                        if not ifLineIsEmptyOrPureComment(tmpLine):
                            if tmpLine[0] != '|':
                                tmpText.append(tmpLine)
                        # print(tmpLine)

            if label != '':
                label2EngTextDict[label] = tmpText

    dlabelDict = dict()

    ws = wb['标']
    for wbRow in range(2, ws.max_row + 1):
        olabel = ws.cell(row = wbRow, column = 4).value
        if olabel in label2EngTextDict:
            dlabel = ws.cell(row = wbRow, column = 3).value
            dmap = ws.cell(row = wbRow, column = 1).value
            oeom_jp = ws.cell(row = wbRow, column = 6).value
            dlabelDict[dlabel] = (label2EngTextDict[olabel],dmap,olabel,oeom_jp)
        else:
            print(f'ERROR: {olabel} not Found!')
    return dlabelDict

def isATextLabel(text):
    components = text.strip().split(' ')
    if len(components) >= 1:
        if 'text' in components[0].lower():
            if components[0][-1] == ':':
                return True
    else:
        return False
    return False

def isALabel(text):
    components = text.strip().split(' ')
    if len(components) >= 1:
        if len(components[0]) < 1:
            return False
        if components[0][-1] == ':':
            return True
    else:
        return False
    return False

def getStrippedLabel(text):
    components = text.strip().split(' ')
    return components[0].replace(':','')

def readLines(path):
    with open(path, 'r') as file:
        lines = file.readlines()
        return lines

def ifLineIsEmptyOrPureComment(line):
    if len(line.strip()) > 0:
        return line.strip()[0] == ';'
    else:
        return True
    
def getGSDict():
    dlabelDict = dict()
    # for file in files:
    pathLines = readLines('xlsx/textScripts.txt')
    

    for pathWithLine in pathLines:
        path = pathWithLine.strip()
        lines = readLines(path)
        tmpText = []
        tmpCmd = []
        label = ''
        foundFirst = False

        for index, line in enumerate(lines):

            if isATextLabel(line):
                nextLine = lines[index + 1]
                if "text_" in nextLine.lower() or '\"' in nextLine or ifLineIsEmptyOrPureComment(nextLine):
                    if (foundFirst) and label != '':
                        dlabelDict[label] = (tmpText,path,"N/A",tmpCmd)
                        # print(path)
                        # print(label)
                        # print(tmpText)
                        # print()

                    foundFirst = True
                    label = getStrippedLabel(line)
                    tmpText = []
                    tmpCmd = []

            elif isALabel(line):
                if (foundFirst) and label != '':
                    dlabelDict[label] = (tmpText,path,"N/A",tmpCmd)
                    # print(path)
                    # print(label)
                    # print(tmpText)
                    # print()

                label = ""
                tmpText = []
                tmpCmd = []
            elif index == len(lines) - 1:
                if (foundFirst) and label != '':
                    dlabelDict[label] = (tmpText,path,"N/A",tmpCmd)
                    # print(path)
                    # print(label)
                    # print(tmpText)
                    # print()


            elif foundFirst:
                
                tmpLine = line.replace(' \"',part).replace('\"',part).replace('\t','').replace('\n','')
                # print(tmpLine)
                components = tmpLine.split(part)
                # print(components)
                if len(components) >= 2:
                    if '\"' in line.strip():
                        tmpText.append(components[1].replace('\"',''))
                        tmpCmd.append(components[0])
                else:
                    components2 = line.strip().split(' ')
                    if len(components2) >= 2 and len(line.strip()) > 0:
                        if line.strip()[0] != ';':
                            newLine = ('+' + line.strip() + '+').strip() 
                            # print(newLine)
                            # print(list(newLine))
                            tmpText.append(newLine)
                            tmpCmd.append(newLine)
                    if tmpLine.strip() == 'done' or tmpLine.strip() == 'prompt' or tmpLine.strip() == 'text_end':
                        tmpCmd.append(tmpLine.strip())
    return dlabelDict

def loopRemove(text):
    line = text
    while '+ ' in line:
        line = line.replace('+ ','+')
    return line

def compareText(d1Lines,d2Lines,mode):
    if mode == 0:
        d1Text = loopRemove((' '.join(d1Lines)).replace('@','').strip().replace(' +','+'))
        d2Text = loopRemove((' '.join(d2Lines)).replace('@','').strip().replace('<PLAY_G>','<PLAYER>').replace('$','¥'))
        return (d1Text.lower() == d2Text.lower(),d1Text,d2Text)
    else:
        d1Text = loopRemove((' '.join(d1Lines)).replace('@','').strip().replace('<PLAY_G>','<PLAYER>').replace('$','¥'))
        d2Text = loopRemove((' '.join(d2Lines)).replace('@','').strip().replace(' +','+'))
        return (d1Text.lower() == d2Text.lower(),d1Text,d2Text)


def findDict1InDict2(dict1,dict2,dict1Names,dict2Names,mode):
    labelNotFound = []
    textNotMatch = []
    filepathNotMatch = []

    FileNotFoundDict = dict()
    for key in dict1:
        if (mode == 1):
            if os.path.exists(RN(dict1[key][1])):
                if RN(dict1[key][1]) != '':
                    FileNotFoundDict[dict1[key][1]] = True

        if not key in dict2:
            txt = (' '.join(dict1[key][0])).replace('@','').strip()
            labelNotFound.append((key,dict1,dict2,dict1Names,dict2Names,txt,''))
            
        else:
            compare = compareText(dict1[key][0],dict2[key][0],mode)
            if not compare[0]:
                textNotMatch.append((key,dict1,dict2,dict1Names,dict2Names,compare[1],compare[2]))
            if dict1[key][1] != dict2[key][1]:
                filepathNotMatch.append((key,dict1,dict2,dict1Names,dict2Names,dict1[key][1],dict2[key][1],compare[1],compare[2]))

    if (mode == 1):
        for key2 in FileNotFoundDict:
            print(key2)

    return [labelNotFound,textNotMatch,filepathNotMatch]


def readTextFile(path):
    file = open(path,"r")
    return file.read()

def generateText(results):
    wb = openpyxl.Workbook()
    ws1 = wb.worksheets[0]
    ws2 = wb.create_sheet("Sheet2")
    rowID1 = 1
    rowID2 = 1
    errorID = 0
    
    html_content = readTextFile('goldJPNtext.html')

    # Parse the content using BeautifulSoup
    soup = BeautifulSoup(html_content, 'html.parser')

    # Extract all 'file' elements
    files = soup.find_all('file')

    JPNDict = dict()
    for file in files:
        # print(f"File Name: {file['name']}")
        messages = file.find_all('message')
        for message in messages:
            pokescript = message.pokescript
        # Extract and print Japanese content
        if pokescript.jpn is not None:
            japanese_text = pokescript.jpn.text.strip()  # Strip to remove extra whitespace
            JPNDict[message['id']] = (japanese_text, pokescript['scriptid'])
        else:
            JPNDict[message['id']] = ('', pokescript['scriptid'])

    for index, errors in enumerate(results):
        for error in errors:
            if index == 1:
                ws1.cell(row = rowID1, column = 1).value = error[1][error[0]][1]
                ws1.cell(row = rowID1, column = 3).value = error[0]
                ws1.cell(row = rowID1, column = 4).value = error[2][error[0]][2]
                ws1.cell(row = rowID1, column = 9).value = 'GS'
                
                

                ws2.cell(row = rowID2, column = 1).value = '|---英文---------+'
                ws2.cell(row = rowID2, column = 3).value = '|---日文-------------------------+'
                ws2.cell(row = rowID2, column = 5).value = '|---翻译--------------+----------+'
                ws2.cell(row = rowID2, column = 8).value = '|---提词---------+'
                ws2.cell(row = rowID2, column = 9).value = '|---注释---------+'
                ws2.cell(row = rowID2, column = 10).value = error[2][error[0]][2]
                ws2.cell(row = rowID2, column = 11).value = 'GS'
                rowID2 += 1
                i = 0

                JPNRowID2 = rowID2
                while i < len(error[1][error[0]][0]):
                    if error[1][error[0]][3][i] == 'para':
                        rowID2 += 1
                    ws2.cell(row = rowID2 + i, column = 1).value = error[1][error[0]][0][i]
                    i += 1


                if error[2][error[0]][2] in JPNDict:
                    jpntextComponents = JPNDict[error[2][error[0]][2]][0].split('\n')
                    for w in range(len(jpntextComponents)):
                        ws2.cell(row = JPNRowID2 + w, column = 3).value = jpntextComponents[w]


                if i < len(error[1][error[0]][3]):
                    ws1.cell(row = rowID1, column = 6).value = error[1][error[0]][3][i]
                    ws1.cell(row = rowID1, column = 7).value = error[1][error[0]][3][i]

                rowID1 += 1
                rowID2 += 1 + i
                errorID += 1

    wb.save('label.xlsx')


def printMatchResults(results):
    generateText(results)
    for index, errors in enumerate(results):
        for error in errors:
            if index == 0:
                print(bcolors.FAIL)
                print(f'FAIL: LabelNotFound in {error[1][error[0]][1]}')
                print(f'key: {error[0]} {error[1][error[0]][2]}')
                print(error[1][error[0]][0])
                print(error[1][error[0]][3])
                print(f'text: {error[5]}')
                print(f'not found in {error[4]}')
            elif index == 1:
                print(bcolors.OKBLUE)
                print(f'FAIL: TextNotMatch in {error[1][error[0]][1]}')
                print(f'key: {error[0]} {error[1][error[0]][2]}')
                print(f'key: {error[0]} {error[2][error[0]][2]}')
                print(f'{error[5]}')
                print(f'{error[6]}')
                # print(f'{list(error[5])}')
                # print(f'{list(error[6])}')
            elif index == 2:
                print(bcolors.ENDC)
                print(f'FAIL: PathNotMatch in {error[1][error[0]][1]}')
                print(f'key: {error[0]} {error[1][error[0]][2]}')
                print(f'key: {error[0]} {error[2][error[0]][2]}')
                print(f'{error[5]}')
                print(f'{error[6]}')
                print(f'{error[7]}')
                print(f'{error[8]}')
                if error[7] == error[8]:
                    print(f'{error[5]},{error[0]},{error[2][error[0]][2]},{error[2][error[0]][3]}')
                    print('Text MATCH')
                else:
                    print('Text NOT Match')

        print(bcolors.OKGREEN)
        


print(bcolors.OKGREEN)
CRDict = getCrystalDict()
GSDict = getGSDict()


results = findDict1InDict2(GSDict,CRDict,"GS Text","Crystal Text",0)
printMatchResults(results)

# results2 = findDict1InDict2(CRDict,GSDict,"Crystal Text","GS Text",1)
# printMatchResults(results2)
# print(GSDict)